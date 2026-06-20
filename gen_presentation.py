#!/usr/bin/env python3
"""Generuje krótką prezentację PPT: LoRa Zigbee SCADA v38 — status + procedury testowe
wszystkich funkcjonalności. Stylistyka spójna z dashboardem (dark #0a0a0a, cyjan #22d3ee)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

BG    = RGBColor(0x0a, 0x0a, 0x0a)
CARD  = RGBColor(0x14, 0x14, 0x14)
CYAN  = RGBColor(0x22, 0xd3, 0xee)
GREEN = RGBColor(0x4a, 0xde, 0x80)
AMBER = RGBColor(0xfb, 0xbf, 0x24)
RED   = RGBColor(0xf8, 0x71, 0x71)
WHITE = RGBColor(0xe5, 0xe5, 0xe5)
GREY  = RGBColor(0x73, 0x73, 0x73)
DGREY = RGBColor(0x52, 0x52, 0x52)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(1, 0, 0, SW, SH)
    bg.fill.solid(); bg.fill.fore_color.rgb = BG; bg.line.fill.background()
    bg.shadow.inherit = False
    return s


def box(s, x, y, w, h, fill=None, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh


def txt(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, sp=1.0):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.02)
    first = True
    for line in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align; p.line_spacing = sp; p.space_after = Pt(2)
        for (t, sz, col, bold) in line:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.color.rgb = col; r.font.bold = bold
            r.font.name = "Consolas"
    return tb


def header(s, kicker, title):
    box(s, 0.5, 0.45, 0.12, 0.55, fill=CYAN)
    txt(s, 0.75, 0.42, 12, 0.35, [[(kicker, 11, DGREY, True)]])
    txt(s, 0.73, 0.66, 12.2, 0.7, [[(title, 26, WHITE, True)]])


def footer(s, n):
    txt(s, 0.75, 7.0, 8, 0.3, [[("LoRa Zigbee SCADA v38 — BMS rozproszony", 9, DGREY, False)]])
    txt(s, 11.8, 7.0, 1.2, 0.3, [[(str(n), 9, DGREY, True)]], align=PP_ALIGN.RIGHT)


def proc_slide(n, kicker, title, steps, expect, status):
    """Slajd procedury testowej: kroki (lewa) + oczekiwany wynik (prawa) + pasek statusu."""
    s = slide(); header(s, kicker, title)
    # status pill
    scol = {"OK": GREEN, "WARN": AMBER, "BLOK": RED}[status[0]]
    box(s, 10.4, 0.5, 2.45, 0.5, fill=CARD, line=scol)
    txt(s, 10.4, 0.52, 2.45, 0.46, [[("● " + status[1], 11, scol, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    # kroki
    box(s, 0.75, 1.55, 7.2, 5.1, fill=CARD, line=RGBColor(0x1f,0x1f,0x1f))
    txt(s, 1.0, 1.7, 6.8, 0.4, [[("PROCEDURA TESTOWA", 12, CYAN, True)]])
    rows = []
    for i, st in enumerate(steps, 1):
        rows.append([(f"{i}. ", 12, CYAN, True), (st, 12, WHITE, False)])
    txt(s, 1.0, 2.15, 6.8, 4.4, rows, sp=1.15)
    # oczekiwany wynik
    box(s, 8.15, 1.55, 4.45, 5.1, fill=CARD, line=RGBColor(0x1f,0x1f,0x1f))
    txt(s, 8.4, 1.7, 4.0, 0.4, [[("OCZEKIWANY WYNIK", 12, GREEN, True)]])
    erows = [[("✓ ", 11, GREEN, True), (e, 11, WHITE, False)] for e in expect]
    txt(s, 8.4, 2.15, 4.0, 4.4, erows, sp=1.15)
    footer(s, n)


# ── 1. TYTUŁ ──
s = slide()
box(s, 0, 0, SW, SH, fill=BG)
box(s, 0.9, 2.5, 0.16, 2.0, fill=CYAN)
txt(s, 1.3, 2.4, 11, 0.5, [[("LoRa ZIGBEE SCADA — v38", 15, CYAN, True)]])
txt(s, 1.28, 2.95, 11.5, 1.4, [[("Rozproszony BMS", 46, WHITE, True)],
                                [("Zigbee → LoRa → Supervisor → Home Assistant", 20, GREY, False)]])
txt(s, 1.3, 5.0, 11, 0.8, [[("Status systemu i procedury testowe wszystkich funkcjonalności", 16, WHITE, False)],
                            [("3 bramki × ~100 urządzeń · 868 MHz · EU duty 10%", 13, DGREY, False)]])
txt(s, 1.3, 6.6, 11, 0.4, [[("2026-06-21  ·  Step 1–5 zaimplementowane", 11, DGREY, True)]])

# ── 2. ARCHITEKTURA ──
s = slide(); header(s, "PRZEGLĄD", "Architektura systemu")
def node(x, y, w, h, title, sub, col):
    box(s, x, y, w, h, fill=CARD, line=col)
    txt(s, x, y+0.15, w, 0.5, [[(title, 15, col, True)]], align=PP_ALIGN.CENTER)
    txt(s, x, y+0.62, w, h-0.6, [[(l, 10, WHITE, False)] for l in sub], align=PP_ALIGN.CENTER, sp=1.05)
node(4.7, 1.7, 3.9, 1.25, "SUPERVISOR (G0)", ["Xeon · HA Supervised", "3 anteny LoRa · agregacja", "calendar.lora_global (master)"], CYAN)
txt(s, 4.7, 3.05, 3.9, 0.4, [[("▲  LoRa 868 MHz · 220B · 3s cooldown  ▼", 11, AMBER, True)]], align=PP_ALIGN.CENTER)
for i,(x,nm) in enumerate([(0.95,"BRAMKA G1"),(4.9,"BRAMKA G2"),(8.85,"BRAMKA G3")]):
    col = GREEN if i==0 else DGREY
    node(x+0.0, 3.6, 3.5, 1.5, nm, ["Z2M + ConBee/SONOFF","HA lokalne · ~100 dev","Heltec V3 (SX1262)"], col)
txt(s, 0.95, 5.3, 11.5, 1.4, [
    [("G1 = AKTYWNA (live test bed: 100.98.155.78)   ·   G0 = 100.79.111.24", 12, GREEN, True)],
    [("Filozofia: każda bramka autonomiczna (lokalne Z2M+HA), supervisor agreguje, minimalizacja ruchu LoRa.", 12, WHITE, False)],
    [("Constraint: 220B max payload, 3.0s+jitter cooldown, half-duplex, EU duty cycle 10%.", 11, GREY, False)],
], sp=1.2)
footer(s, 2)

# ── 3. STATUS GOTOWOŚCI ──
s = slide(); header(s, "PODSUMOWANIE", "Status gotowości do wdrożenia")
box(s, 0.75, 1.55, 6.0, 5.1, fill=CARD, line=RGBColor(0x1f,0x3a,0x1f))
txt(s, 1.0, 1.7, 5.5, 0.4, [[("DZIAŁA (zweryfikowane live)", 13, GREEN, True)]])
ok = ["Detekcja offline/online (z2m → bramka)","Sync availability bramka → supervisor (spójny)",
      "Anomalie: offline / bateria / stagnacja","Popup z listą + clear per-urządzenie",
      "Kalendarz PUSH (supervisor → bramka)","Tryb pracy bramki (day/night/all-time)",
      "Dashboard bramki: Harmonogram/Pomiary/Alarmy","LQI / link quality (force_update)",
      "Sterowanie switch + cmd retry→offline","Parametry T1/T2/T3 + Send Config"]
txt(s, 1.0, 2.2, 5.5, 4.4, [[("✓ ", 11, GREEN, True),(o, 11, WHITE, False)] for o in ok], sp=1.18)
box(s, 7.0, 1.55, 5.6, 5.1, fill=CARD, line=RGBColor(0x3a,0x2f,0x1f))
txt(s, 7.25, 1.7, 5.2, 0.4, [[("WYMAGA INTERWENCJI (sprzęt/decyzja)", 13, AMBER, True)]])
todo = [("BLOK","Kalendarz PULL (bramka→sup): chunk 100B w górę ginie — asymetria RF (anteny/TX)"),
        ("WARN","Switche: offline po timeout = fałszywy (nie raportują) → potrzebny z2m active-ping"),
        ("WARN","Commit: config.py ma realne sekrety → scrub przed git"),
        ("INFO","Krok 8 (sterowanie pełne) i Krok 17 (parametry bidir) z planu 17")]
rows=[]
for tag,t in todo:
    col={"BLOK":RED,"WARN":AMBER,"INFO":DGREY}[tag]
    rows.append([(f"[{tag}] ",11,col,True),(t,11,WHITE,False)])
txt(s, 7.25, 2.25, 5.2, 4.3, rows, sp=1.3)
footer(s, 3)

# ── 4..N PROCEDURY ──
proc_slide(4, "FUNKCJA 01", "Detekcja offline / online urządzeń",
    ["Na bramce: odłącz zasilanie urządzenia Zigbee (np. Temp 1) lub wyjmij baterię.",
     "Odczekaj > timeout danego typu (T1 switch 58min / T2 sensor 60min / T3 door-leak 120min).",
     "Obserwuj log bramki: '💀 <dev> OFFLINE (cisza z2m ...)'.",
     "Sprawdź dashboard bramki (zakładka Bramka) — kafel OFFLINE rośnie, urządzenie czerwone.",
     "Sprawdź supervisor: sensor.lora_gw_g1_devices_offline + binary_sensor ...available = off.",
     "Podłącz urządzenie → następny raport z2m → powrót.",],
    ["Bramka i supervisor pokazują TO SAMO urządzenie offline.",
     "Licznik = lista = *_available (spójność).",
     "Po powrocie: '🟢 ONLINE' i licznik spada.",
     "Reconcile co 900s łapie zgubione pakiety."],
    ("OK","DZIAŁA"))

proc_slide(5, "FUNKCJA 02", "Anomalie + popup z listą + clear",
    ["Wywołaj anomalię: offline (jak F01), niska bateria (<25%), lub stagnacja (brak zmian 72h).",
     "Na dashboardzie supervisora (widok Bramki) kliknij kafel OFFLINE / BATERIA / INNE.",
     "Otworzy się popup z dynamiczną listą: urządzenie · bramka · typ · data godzina.",
     "Kliknij wiersz konkretnej anomalii → clear tej jednej anomalii.",
     "Lub użyj 'WYCZYŚĆ WSZYSTKIE <typ>' (z potwierdzeniem).",],
    ["Popup pokazuje listę z timestampem DD.MM HH:MM.",
     "Klik wiersza → '🗑️ clear ręczny ... n=1' w logu.",
     "Realne offline wraca po reconcile (poprawne).",
     "Każda anomalia z nazwą + ID bramki."],
    ("OK","DZIAŁA"))

proc_slide(6, "FUNKCJA 03", "Kalendarz PUSH (supervisor → bramka)",
    ["Na supervisorze edytuj calendar.lora_global — dodaj wydarzenie (PRODUKCJA/PRZERWA/SERWIS).",
     "Naciśnij przycisk 'Sync Calendar All' (lub per-bramka 'Sync Calendar').",
     "Obserwuj log supervisora: 'HA → N slotów' → cal_begin/chunk/end.",
     "Obserwuj log bramki: '📥 Odebrano N slotów' → 'ICS zapisany' → 'HA reload'.",
     "Sprawdź calendar.lora_g1 na bramce — ma nowe wydarzenie.",
     "Sprawdź sensor ...tryb i ...hash_kalendarza (zmiana).",],
    ["calendar.lora_g1 = calendar.lora_global (forward).",
     "Hash kalendarza zmienia się i zgadza po obu str.",
     "Tryb produkcji aktualny + następna zmiana.",
     "Drift-hash w HB auto-resyncuje.",
     "Transfer z CRC + ACK + retransmit."],
    ("OK","DZIAŁA"))

proc_slide(7, "FUNKCJA 04", "Kalendarz PULL (bramka → supervisor)",
    ["Na bramce edytuj harmonogram lokalnie (wymaga gw_calendar_id w config).",
     "Naciśnij 'Push Schedule Up' na dashboardzie bramki.",
     "Obserwuj log bramki: 'Begin gw_push' → cal_chunk.",
     "Obserwuj supervisor: 'Begin gw_push' → 'Odebrano' → mirror calendar.lora_g1.",
     "BLOKER: chunk 100B w górę chronicznie ginie (RF asymetria) → 'przerwane'.",],
    ["Mechanizm wired: begin dociera, retransmit działa.",
     "Mirror-only (bez pętli, master = global).",
     "⚠ Wymaga poprawy RF w górę (sprzęt):",
     "↑TX bramki / antena RX sup / mniejszy chunk."],
    ("BLOK","BLOK RF"))

proc_slide(8, "FUNKCJA 05", "Tryb pracy bramki (day / night / all-time)",
    ["W config bramki ustaw gateway_mode.operating_mode = 'day' lub 'night' (+ lat/lon lub godziny).",
     "Restart procesu bramki (przez Launcher).",
     "Sprawdź sensor ...tryb_pracy_bramki = Dzienna/Nocna/Całodobowa.",
     "W oknie NIEAKTYWNYM: urządzenia bez zasilania NIE są zgłaszane jako offline (supresja).",
     "Sprawdź pole 'ga' w HB (1=aktywna, 0=wstrzymana) i kartę na dashboardzie.",],
    ["Tryb widoczny jako encja + karta (Harmonogram).",
     "Supresja offline gdy bramka nieaktywna.",
     "Solar (sunrise equation) lub stałe godziny.",
     "all-time = brak supresji (domyślnie)."],
    ("OK","DZIAŁA*"))

proc_slide(9, "FUNKCJA 06", "Sterowanie + dashboardy + LQI",
    ["Dashboard bramki → zakładka Sterowanie → kliknij switch (Test 1/2) → ON/OFF.",
     "Obserwuj: komenda → Z2M → potwierdzenie 'st' → odbicie stanu w HA.",
     "Brak potwierdzenia: retry ×2 (15s) → '💀 brak st → OFFLINE'.",
     "Zakładka Pomiary: temp/wilgotność + wykresy 24h; Alarmy: leak/door.",
     "Każdy kafel: ● online/offline · 🔋 bateria · 📶 LQI (link quality).",],
    ["Switch sterowalny dwukierunkowo (tap→toggle).",
     "Zmiana zewnętrzna też odbita (propagacja z2m).",
     "LQI świeży (force_update co raport).",
     "Ikony w osobnej kolumnie (czytelne)."],
    ("OK","DZIAŁA"))

# ── N+1: TODO USERA ──
s = slide(); header(s, "DLA CIEBIE", "Co musisz zrobić Ty (ręce / decyzje)")
items = [
    (RED,  "1. RF w górę (PULL kalendarza)", "Przy maszynach: zwiększ TX power bramki LUB lepsza antena RX supervisora dla G1. Bez tego dwukierunkowy kalendarz (Push Schedule Up) nie domknie się — chunk 100B ginie."),
    (AMBER,"2. z2m active availability (switche)", "Na bramce w z2m configuration.yaml: 'availability:' + 'advanced.last_seen: ISO_8601' → restart z2m. Naprawia fałszywe offline switchy (nie raportują = wyglądają martwe)."),
    (AMBER,"3. Scrub sekretów przed commitem", "config.py ma realne MQTT_PASS + tokeny. Uruchom scrub_secrets.py / _cfg_scrub.py LUB 'git rm --cached config.py' + dopisz do .gitignore, potem commit."),
    (DGREY,"4. Decyzja: kroki 8 i 17", "Czy budujemy pełne sterowanie (krok 8) i parametry bidirectional (krok 17) z planu 17-kroków? — nowa praca, nie bug."),
    (DGREY,"5. Restart Claude Code (opcjonalnie)", "Załaduje tokeny MCP ha-gw/ha-sup (teraz pracuję przez REST/SSH — działa, ale MCP byłoby wygodniejsze)."),
]
y = 1.6
for col, t, d in items:
    box(s, 0.75, y, 0.12, 0.95, fill=col)
    txt(s, 1.0, y, 11.6, 0.45, [[(t, 14, col if col!=DGREY else WHITE, True)]])
    txt(s, 1.0, y+0.42, 11.6, 0.55, [[(d, 11, GREY, False)]], sp=1.05)
    y += 1.05
footer(s, 10)

# ── N+2: WDROŻENIE / PODSUMOWANIE ──
s = slide(); header(s, "WDROŻENIE", "Gotowość i procedura wdrożenia")
box(s, 0.75, 1.6, 5.7, 5.0, fill=CARD, line=RGBColor(0x1f,0x1f,0x1f))
txt(s, 1.0, 1.75, 5.2, 0.4, [[("OCENA GOTOWOŚCI", 13, CYAN, True)]])
txt(s, 1.0, 2.25, 5.2, 4.2, [
    [("Rdzeń (monitoring + anomalie + ",12,WHITE,False),("PUSH",12,GREEN,True),(")",12,WHITE,False)],
    [("   → GOTOWY do wdrożenia produkcyjnego.",12,GREEN,True)],
    [("",6,WHITE,False)],
    [("Dwukierunkowy kalendarz (PULL)",12,WHITE,False)],
    [("   → po poprawie RF (sprzęt).",12,AMBER,True)],
    [("",6,WHITE,False)],
    [("Pełna detekcja switchy",12,WHITE,False)],
    [("   → po z2m availability (faza 2).",12,AMBER,True)],
    [("",6,WHITE,False)],
    [("Werdykt: rdzeń production-ready;",12,WHITE,True)],
    [("2 pozycje sprzętowe do domknięcia",12,WHITE,True)],
    [("pełni dwukierunkowej.",12,WHITE,True)],
], sp=1.15)
box(s, 6.65, 1.6, 5.95, 5.0, fill=CARD, line=RGBColor(0x1f,0x1f,0x1f))
txt(s, 6.9, 1.75, 5.5, 0.4, [[("PROCEDURA WDROŻENIA (per bramka)", 13, CYAN, True)]])
dep = ["Postaw Heltec V3 + ConBee/SONOFF + Z2M + HA lokalne.",
       "Ustaw .role (gateway/supervisor) w ~/meshtastic/.",
       "Wpisz sekrety in-place (MQTT_PASS + HA_TOKEN) — NIE scp configu.",
       "Skonfiguruj monitored[] / priority[] / T1-T3 / kalendarz.",
       "Start przez Launcher (:8765) → test_step5_anomaly.py.",
       "Wygeneruj dashboard: gen_gateway_dashboard.py + patche sup.",
       "Smoke: ping_all → HB → discovery → batch → anomalia testowa.",
       "Walidacja: offline-sync + PUSH kalendarza + sterowanie."]
txt(s, 6.9, 2.25, 5.5, 4.2, [[(f"{i}. ",11,CYAN,True),(d,11,WHITE,False)] for i,d in enumerate(dep,1)], sp=1.2)
footer(s, 11)

import os
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "SCADA_v38_test_procedures.pptx")
prs.save(out)
print("zapisano:", out, "| slajdów:", len(prs.slides._sldIdLst))
